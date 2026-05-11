from flask import Flask, render_template, request, redirect, url_for, session
from flask_session import Session
import pandas as pd
import calendar
import math
import os
import datetime

app = Flask(__name__)
app.secret_key = 'sirekap_secret_key'

# ================= CONFIG =================
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['SESSION_TYPE'] = 'filesystem'
app.config['SESSION_PERMANENT'] = False

Session(app)

# ================= LOGIN =================
@app.route('/', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')

        if username == "admin" and password == "sirekap":
            session.clear()
            session['user'] = username
            return redirect(url_for('dashboard'))

    return render_template('login.html')

# ================= LOAD DATA =================
def load_data(file):
    df = pd.read_excel(file, dtype=str)

    df.columns = df.columns.str.strip().str.lower()

    required = ['id number', 'name', 'date/time']
    for col in required:
        if col not in df.columns:
            raise ValueError(f"Kolom '{col}' tidak ditemukan")

    df['id number'] = df['id number'].fillna("").astype(str).str.strip()
    df['name'] = df['name'].fillna("").astype(str).str.strip()

    df['date/time'] = pd.to_datetime(df['date/time'], dayfirst=True, errors='coerce')
    df = df.dropna(subset=['date/time'])

    df['tanggal'] = df['date/time'].dt.date

    return df

# ================= CREATE ATTENDANCE =================
def create_attendance(df, bulan, tahun, minggu="all"):

    pegawai = df[['id number', 'name']].drop_duplicates().reset_index(drop=True)

    last_day = calendar.monthrange(tahun, bulan)[1]

    tanggal_range = pd.date_range(
        start=f"{tahun}-{bulan:02d}-01",
        end=f"{tahun}-{bulan:02d}-{last_day}",
        freq='D'
    )

    if minggu != "all":
        minggu = int(minggu)
        tanggal_range = [tgl for tgl in tanggal_range if math.ceil(tgl.day / 7) == minggu]

    tabel = pd.DataFrame()

    tabel['NIP'] = pegawai['id number'].fillna("").astype(str).str.strip()
    tabel['Nama'] = pegawai['name'].fillna("").astype(str).str.strip()

    tabel['NIP'] = tabel['NIP'].replace(["", "nan", "None"], "-")

    tanggal_cols = []
    day_map = {}

    for tgl in tanggal_range:
        col = str(tgl.day)
        tanggal_cols.append(col)

        tabel[col] = ""

        hadir = df[df['tanggal'] == tgl.date()]['id number'].astype(str).str.strip().unique()

        for i in tabel.index:
            if tabel.loc[i, 'NIP'] in hadir:
                tabel.loc[i, col] = "H"

        if tgl.weekday() == 5:
            day_map[col] = "sabtu"

        elif tgl.weekday() == 6:
            day_map[col] = "minggu"

        else:
            day_map[col] = ""

    tabel['Total'] = tabel[tanggal_cols].apply(
        lambda x: sum(1 for v in x if v == "H"),
        axis=1
    )

    tabel.insert(0, "No", range(1, len(tabel) + 1))
    tabel = tabel[['No', 'NIP', 'Nama'] + tanggal_cols + ['Total']]

    return tabel, day_map

# ================= DASHBOARD =================
@app.route('/dashboard', methods=['GET', 'POST'])
def dashboard():
    if 'user' not in session:
        return redirect(url_for('login'))

    tabel = session.get('tabel', [])
    kolom = session.get('kolom', [])
    day_map = session.get('day_map', {})

    if request.method == 'POST':
        aksi = request.form.get('aksi')
        print("AKSI:", aksi)

        try:
            # ================= LOAD =================
            if aksi == "load":
                file = request.files.get('file')
                bulan = request.form.get('bulan')
                tahun = request.form.get('tahun')

                if file and bulan and tahun:
                    filepath = os.path.join(UPLOAD_FOLDER, file.filename)
                    file.save(filepath)

                    session['filename'] = file.filename

                    df = load_data(filepath)
                    df_tabel, day_map = create_attendance(df, int(bulan), int(tahun))

                    tabel = df_tabel.to_dict(orient='records')
                    kolom = list(df_tabel.columns)

                    session['tabel'] = tabel
                    session['kolom'] = kolom
                    session['day_map'] = day_map
                    session['tabel_awal'] = [row.copy() for row in tabel]
                    session['history'] = []
                    session['redo'] = []
                    session['riwayat'] = []

                    session.modified = True

            # ================= UPDATE =================
            elif aksi == "update":
                tabel = session.get('tabel', [])
                kolom = session.get('kolom', [])
                day_map = session.get('day_map', {})

                if not tabel:
                    return "ERROR: Load data dulu"

                nama_list = request.form.getlist('nama[]')
                mulai_list = request.form.getlist('mulai[]')
                selesai_list = request.form.getlist('selesai[]')
                status_list = request.form.getlist('status[]')

                session.setdefault('history', []).append([row.copy() for row in tabel])
                session['redo'] = []

                riwayat = session.get('riwayat', [])

                for i in range(len(nama_list)):
                    nama = nama_list[i]
                    mulai = mulai_list[i]
                    selesai = selesai_list[i]
                    status = status_list[i]

                    if not nama or not mulai or not selesai or not status:
                        continue

                    mulai = int(mulai)
                    selesai = int(selesai)

                    for row in tabel:
                        if row['Nama'] == nama:
                            for t in range(mulai, selesai + 1):
                                t = str(t)
                                if t in row:
                                    row[t] = status

                    riwayat.append({
                        "nama": nama,
                        "mulai": mulai,
                        "selesai": selesai,
                        "status": status,
                        "waktu": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    })

                # HITUNG ULANG TOTAL
                tanggal_cols = [k for k in kolom if k in day_map]

                for row in tabel:
                    row['Total'] = sum(1 for k in tanggal_cols if row.get(k) == "H")

                # BATASI RIWAYAT
                if len(riwayat) > 100:
                    riwayat = riwayat[-100:]

                session['tabel'] = tabel
                session['riwayat'] = riwayat
                session.modified = True

            # ================= UNDO =================
            elif aksi == "undo":
                history = session.get('history', [])
                if history:
                    session.setdefault('redo', []).append(session['tabel'])
                    session['tabel'] = history.pop()
                    session['history'] = history
                    session.modified = True

            # ================= REDO =================
            elif aksi == "redo":
                redo = session.get('redo', [])
                if redo:
                    session.setdefault('history', []).append(session['tabel'])
                    session['tabel'] = redo.pop()
                    session['redo'] = redo
                    session.modified = True

            # ================= RESET =================
            elif aksi == "reset":
                session['tabel'] = session.get('tabel_awal', [])
                session['history'] = []
                session['redo'] = []
                session['riwayat'] = []
                session.modified = True

        except Exception as e:
            return f"ERROR: {str(e)}"

    return render_template(
        'dashboard.html',
        tabel=session.get('tabel', []),
        kolom=session.get('kolom', []),
        day_map=session.get('day_map', {}),
        filename=session.get('filename'),
        riwayat=session.get('riwayat', [])
    )

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from flask import send_file
import io

@app.route('/download')
def download_excel():
    tabel = session.get('tabel', [])
    kolom = session.get('kolom', [])
    day_map = session.get('day_map', {})

    if not tabel:
        return "Tidak ada data"

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Absensi"

    # ================= WARNA =================
    header_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    sabtu_fill = PatternFill(start_color="053DB3", end_color="053DB3", fill_type="solid")   # biru
    minggu_fill = PatternFill(start_color="E60008", end_color="E60008", fill_type="solid")  # merah
    dl_fill = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")      # oranye

    # ================= BORDER =================
    thin = Side(style='thin')
    thick = Side(style='medium')

    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_header = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thick')
    )

    # ================= HEADER =================
    ws.append(kolom)

    for i, cell in enumerate(ws[1]):
        cell.font = Font(bold=True, color="000000", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill
        cell.border = border_header

        # 🔥 warna sabtu minggu di header
        col_name = kolom[i]
        if col_name in day_map:
            if day_map[col_name] == "sabtu":
                cell.fill = sabtu_fill
            elif day_map[col_name] == "minggu":
                cell.fill = minggu_fill

    # ================= ISI DATA =================
    for row_data in tabel:
        row = [row_data.get(k, "") for k in kolom]
        ws.append(row)

    # ================= STYLE ISI =================
    for row in ws.iter_rows(min_row=2):
        for j, cell in enumerate(row):
            val = str(cell.value).strip()

            # border
            cell.border = border_all

            # alignment
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # DL warna oranye
            if val == "DL":
                cell.fill = dl_fill

            # sabtu minggu kolom
            col_name = kolom[j]
            if col_name in day_map:
                if day_map[col_name] == "sabtu":
                    cell.fill = sabtu_fill
                elif day_map[col_name] == "minggu":
                    cell.fill = minggu_fill

    # ================= AUTO WIDTH =================
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max_length + 3

    # ================= FREEZE HEADER =================
    ws.freeze_panes = "A2"

    # ================= SAVE =================
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="rekap_absensi.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# ================= LOGOUT =================
@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ================= RUN =================
if __name__ == "__main__":
    import os

    port = int(os.environ.get("PORT", 5000))

    app.run(
        host="0.0.0.0",
        port=port
    )